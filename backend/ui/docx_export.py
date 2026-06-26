"""
Server-side DOCX generation using python-docx.
Follows the German reference template specs (Lebenslauf_Muster / Anschreiben_Muster).
"""
import base64
import io
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ── Colours ───────────────────────────────────────────────────────────────────
_BODY_COLOR = RGBColor(0x2B, 0x2B, 0x2B)
_HEADING_COLOR = RGBColor(0x1F, 0x4E, 0x79)
_DATE_COLOR = RGBColor(0x6B, 0x6B, 0x6B)

# ── Section alias map: normalise variant section names to canonical ────────────
_SECTION_ALIASES: dict[str, str] = {
    "KURZPROFIL": "PROFIL",
    "PROFIL/ZUSAMMENFASSUNG": "PROFIL",
    "ZUSAMMENFASSUNG": "PROFIL",
    "BERUFSERFAHRUNG": "PRAKTISCHE ERFAHRUNG",
    "ARBEITSERFAHRUNG": "PRAKTISCHE ERFAHRUNG",
    "TECHNISCHE KENNTNISSE": "KENNTNISSE",
    "SPRACHKENNTNISSE": "SPRACHEN",
    "SPRACHKOMPETENZ": "SPRACHEN",
    "FACHKENNTNISSE": "KENNTNISSE",
}

# ── All recognised German CV section names (canonical + variant spellings) ────
# Used as a whitelist to detect plain-text section headings (no ## or **).
_ALL_SECTION_NAMES: frozenset[str] = frozenset({
    "PROFIL", "KURZPROFIL", "PROFIL/ZUSAMMENFASSUNG", "ZUSAMMENFASSUNG",
    "AUSBILDUNG", "STUDIUM",
    "PRAKTISCHE ERFAHRUNG", "BERUFSERFAHRUNG", "ARBEITSERFAHRUNG", "ERFAHRUNG",
    "PRAKTIKA", "PRAKTIKUM",
    "PROJEKTE", "PROJEKT", "HOCHSCHULPROJEKTE",
    "KENNTNISSE", "TECHNISCHE KENNTNISSE", "FACHKENNTNISSE",
    "IT-KENNTNISSE", "EDV-KENNTNISSE", "SOFT SKILLS",
    "SPRACHEN", "SPRACHKENNTNISSE", "SPRACHKOMPETENZ",
    "ZERTIFIKATE", "ZERTIFIZIERUNGEN", "WEITERBILDUNG", "FORTBILDUNG",
    "EHRENAMT", "EHRENAMTLICHES ENGAGEMENT", "ENGAGEMENT",
    "INTERESSEN", "HOBBYS", "HOBBIES",
    "REFERENZEN", "FREIWILLIGENARBEIT",
    "PUBLIKATIONEN", "VERÖFFENTLICHUNGEN",
    "LEISTUNGEN", "AUSZEICHNUNGEN",
})

# ── Regexes used in normalisation ─────────────────────────────────────────────
_HR_RE = re.compile(r'^[-*_]{3,}\s*$')
_MD_LINK_RE = re.compile(r'\[([^\]]*)\]\([^)]*\)')
_HEADING_RE = re.compile(r'^(#{1,2})\s+(.+)$')
# Matches standalone bold-only heading lines: **SECTION NAME** (all uppercase, 3+ chars)
_BOLD_HEADING_RE = re.compile(r'^\*\*([A-ZÄÖÜ][A-ZÄÖÜ\s/]{2,})\*\*\s*$')
# Unescapes Markdown backslash sequences: \+ \- \. \_ etc.
_MD_UNESCAPE_RE = re.compile(r'\\([*_{}\[\]()#+\-.!|\\`>])')
# German date period: "04.2026 – heute", "03.2026 – 04.2026", "2023 – heute", etc.
_GERMAN_DATE_PERIOD_RE = re.compile(
    r'(?:\d{2}\.)?\d{4}\s*[-–—]\s*(?:(?:\d{2}\.)?\d{4}|heute)',
    re.IGNORECASE,
)
# Tech-Stack line: "Tech-Stack:", "Tech Stack:" — omitted from DOCX body.
_TECH_STACK_LINE_RE = re.compile(
    r'^\*{0,2}[Tt]ech[-\s][Ss]tack\s*\*{0,2}:',
    re.IGNORECASE,
)


def _is_plain_section_line(line: str) -> bool:
    """Return True if the stripped line is a known plain-text CV section heading."""
    return line.strip().upper() in _ALL_SECTION_NAMES


def _extract_known_section(line: str) -> str | None:
    """Return canonical section name if line is a known CV section heading; else None.

    Supports all heading formats the LLM may produce:
      ## PROFIL       # PROFIL
      ## **Kurzprofil**   # **Kurzprofil**
      **KURZPROFIL**      plain PROFIL

    Will NOT match document-title lines, contact lines, or entry titles such as:
      # **Javier Briceño Ticona**
      # **Adresse**: Siegen...
      ## **Universität Siegen · Informatik B.Sc.** | 04.2024 – heute
    """
    stripped = line.strip()
    # # or ## heading (with or without ** wrapping the name)
    hm = _HEADING_RE.match(stripped)
    if hm:
        name = hm.group(2).strip().replace("**", "").strip().upper()
        if name in _ALL_SECTION_NAMES or name in _SECTION_ALIASES:
            return _SECTION_ALIASES.get(name, name)
        return None
    # **ALL CAPS** bold-only line
    bm = _BOLD_HEADING_RE.match(stripped)
    if bm:
        name = bm.group(1).strip().upper()
        if name in _ALL_SECTION_NAMES or name in _SECTION_ALIASES:
            return _SECTION_ALIASES.get(name, name)
        return None
    # Plain whitelist section name (PROFIL, AUSBILDUNG, etc.)
    if _is_plain_section_line(stripped):
        name = stripped.upper()
        return _SECTION_ALIASES.get(name, name)
    return None


# ── Markdown pre-processing ───────────────────────────────────────────────────

def _normalize_cv_markdown(markdown: str, candidate_name: str = "") -> str:
    """
    Pre-process tailored CV markdown before DOCX rendering.

    Steps:
    1. Strip the pre-sections header block (everything before the first KNOWN section)
       when a deterministic candidate_name is already being rendered.
       Uses _extract_known_section() so only whitelisted section names are accepted as
       anchors — document title, contact lines, and entry titles are never treated as
       section anchors. Line 0 is always skipped (it is the document title/name line).
    2. Strip horizontal rules (---, ***, ___).
    3. Convert [display](url) markdown links to display text only.
    4. Unescape Markdown backslash sequences (\\+, \\-, \\., \\_, etc.).
    5. Convert standalone **ALL CAPS** bold lines that are known sections → ## CANONICAL.
    6. Normalise # / ## headings:
         - Known section name → ## CANONICAL (always double-hash, normalised)
         - Non-section heading (document title, entry title, contact line) → plain body text
    7. Convert plain whitelist section headings (PROFIL, AUSBILDUNG, …) → ## CANONICAL.
    8. Strip any remaining ** markers from body lines.
    """
    lines = markdown.splitlines()

    # Strip pre-sections header block when we have a deterministic header.
    # _extract_known_section() rejects document titles, contact lines, and entry titles,
    # so only a real CV section name can be the cut point.
    #
    # Skip this entirely if line 0 is itself a known section — the CV starts clean
    # with no personal header block to strip.
    # Otherwise, search from i > 0 so the document-title line (# **Name** or the
    # candidate's name in plain text) is never treated as the stripping anchor.
    if candidate_name:
        first_line_is_section = bool(lines) and _extract_known_section(lines[0]) is not None
        if not first_line_is_section:
            first_section_idx: int | None = next(
                (i for i, ln in enumerate(lines)
                 if i > 0 and _extract_known_section(ln) is not None),
                None,
            )
            if first_section_idx is not None:
                lines = lines[first_section_idx:]

    out: list[str] = []
    for line in lines:
        stripped = line.strip()

        # Drop horizontal rules
        if _HR_RE.match(stripped):
            continue

        # Convert markdown links → display text
        line = _MD_LINK_RE.sub(r'\1', line)

        # Unescape Markdown backslash sequences: \+ \- \. \_ etc.
        line = _MD_UNESCAPE_RE.sub(r'\1', line)
        stripped = line.strip()

        # Convert standalone **ALL CAPS** bold heading lines to ## CANONICAL.
        # Only proceeds if the content is a known section name.
        bm = _BOLD_HEADING_RE.match(stripped)
        if bm:
            name = bm.group(1).strip().upper()
            canonical = _SECTION_ALIASES.get(name, name)
            if canonical in _ALL_SECTION_NAMES or name in _ALL_SECTION_NAMES:
                out.append(f"## {canonical}")
            else:
                out.append(bm.group(1).strip())  # not a section — plain body text
            continue

        # Normalise # / ## headings.
        # Known section name → ## CANONICAL.
        # Everything else (document title, entry title, contact line) → plain body text.
        hm = _HEADING_RE.match(stripped)
        if hm:
            name = hm.group(2).strip().replace("**", "").strip().upper()
            canonical = _SECTION_ALIASES.get(name, name)
            if canonical in _ALL_SECTION_NAMES or name in _ALL_SECTION_NAMES:
                out.append(f"## {canonical}")
            else:
                body = hm.group(2).strip().replace("**", "").strip()
                out.append(body)
            continue

        # Convert plain whitelist section headings: PROFIL, AUSBILDUNG, etc. → ## CANONICAL
        if _is_plain_section_line(stripped):
            name = stripped.upper()
            canonical = _SECTION_ALIASES.get(name, name)
            out.append(f"## {canonical}")
            continue

        # Strip any remaining ** markers from body lines
        line = line.replace("**", "")

        out.append(line)

    return "\n".join(out)


# ── URL / city helpers ────────────────────────────────────────────────────────

def _clean_url(url: str) -> str:
    """Strip https:// or http:// prefix so URLs display cleanly."""
    return re.sub(r'^https?://', '', url).rstrip('/')


def _short_city(city: str) -> str:
    """Return only the city portion, dropping any state/Bundesland after a comma.

    'Siegen, Nordrhein-Westfalen' → 'Siegen'
    'Frankfurt am Main' → 'Frankfurt am Main' (no comma, unchanged)
    """
    return city.split(",")[0].strip() if city else ""


# ── Low-level XML helpers ─────────────────────────────────────────────────────

def _set_para_spacing(para, before_pt: float = 0, after_pt: float = 0) -> None:
    pPr = para._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), str(int(before_pt * 20)))
    spacing.set(qn("w:after"), str(int(after_pt * 20)))
    pPr.append(spacing)


def _set_run_font(run, name: str = "Calibri", size_pt: float | None = None,
                  color: RGBColor | None = None, bold: bool = False,
                  italic: bool = False) -> None:
    run.font.name = name
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if color is not None:
        run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic


def _add_run(para, text: str, size_pt: float | None = None,
             color: RGBColor | None = None, bold: bool = False,
             italic: bool = False) -> None:
    run = para.add_run(text)
    _set_run_font(run, size_pt=size_pt, color=color or _BODY_COLOR,
                  bold=bold, italic=italic)


def _set_tab_stops(para, right_pos_cm: float) -> None:
    # 1 cm = 1440/2.54 ≈ 566.93 twips (OOXML w:pos unit)
    pPr = para._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(int(right_pos_cm * 1440 / 2.54)))
    tabs.append(tab)
    pPr.append(tabs)


# ── Table helpers ─────────────────────────────────────────────────────────────

def _remove_table_borders(table) -> None:
    """Remove all visible borders from a python-docx Table object."""
    tbl_el = table._tbl
    tblPr = tbl_el.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl_el.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        tblBorders.append(el)
    tblPr.append(tblBorders)


# ── Photo cropping ────────────────────────────────────────────────────────────

def _crop_portrait(image_bytes: bytes) -> bytes | None:
    """
    Center-crop and resize image to portrait 3:4 ratio (225×300 px).
    Returns JPEG bytes, or None if processing fails.
    Requires Pillow (already a project dependency via avatar.py).
    """
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(image_bytes)).convert("RGB")
        w, h = img.size
        target_ratio = 3.0 / 4.0  # width / height

        current_ratio = w / h
        if abs(current_ratio - target_ratio) > 0.02:
            if current_ratio > target_ratio:
                # Too wide — crop width to portrait ratio
                new_w = int(h * target_ratio)
                x_off = (w - new_w) // 2
                img = img.crop((x_off, 0, x_off + new_w, h))
            else:
                # Too tall — crop height to portrait ratio
                new_h = int(w / target_ratio)
                y_off = (h - new_h) // 2
                img = img.crop((0, y_off, w, y_off + new_h))

        img = img.resize((225, 300), Image.LANCZOS)
        out = io.BytesIO()
        img.save(out, format="JPEG", quality=85)
        return out.getvalue()
    except Exception:
        return None


# ── Document setup ────────────────────────────────────────────────────────────

def _setup_a4(doc: Document, top_cm=2.0, bottom_cm=2.0,
              left_cm=2.0, right_cm=2.0, header_dist_cm=1.25) -> None:
    for section in doc.sections:
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = Cm(top_cm)
        section.bottom_margin = Cm(bottom_cm)
        section.left_margin = Cm(left_cm)
        section.right_margin = Cm(right_cm)
        section.header_distance = Cm(header_dist_cm)
        section.footer_distance = Cm(header_dist_cm)


def _set_doc_default_font(doc: Document, name: str = "Calibri",
                          size_pt: float = 11.0) -> None:
    style = doc.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(size_pt)
    style.font.color.rgb = _BODY_COLOR


# ── Shared paragraph helpers ──────────────────────────────────────────────────

def _add_blank_line(doc: Document, size_pt: float = 4.0) -> None:
    para = doc.add_paragraph()
    _set_para_spacing(para, before_pt=0, after_pt=0)
    para.add_run().font.size = Pt(size_pt)


# ── CV-specific helpers ───────────────────────────────────────────────────────

def _add_cv_section_heading(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    _set_para_spacing(para, before_pt=11, after_pt=4)
    run = para.add_run(text.upper())
    _set_run_font(run, size_pt=10.5, color=_HEADING_COLOR, bold=True)
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_cv_entry_title(doc: Document, title: str, org: str = "",
                        period: str = "", usable_width_cm: float = 17.0) -> None:
    para = doc.add_paragraph()
    _set_para_spacing(para, before_pt=10, after_pt=2)
    _add_run(para, title, size_pt=10.5, bold=True)
    if org:
        _add_run(para, f"  |  {org}", size_pt=10.5)
    if period:
        _set_tab_stops(para, right_pos_cm=usable_width_cm)
        para.add_run("\t")
        _add_run(para, period, size_pt=9.5, color=_DATE_COLOR, italic=True)


def _add_cv_bullet(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    _set_para_spacing(para, before_pt=1, after_pt=1)
    pPr = para._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "360")
    ind.set(qn("w:hanging"), "180")
    pPr.append(ind)
    _add_run(para, "–  " + text, size_pt=10.5)


# ── Markdown-inline bold renderer ─────────────────────────────────────────────

def _render_inline(para, text: str, size_pt: float = 10.5,
                   color: RGBColor | None = None) -> None:
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            _add_run(para, part[2:-2], size_pt=size_pt, color=color, bold=True)
        elif part:
            _add_run(para, part, size_pt=size_pt, color=color)


# ── Entry-title heuristics ────────────────────────────────────────────────────

def _split_title_date(line: str) -> tuple[str, str] | None:
    """
    Split an entry title line into (title, date_period).

    Detects German-CV lines of the form:
      "Title | 04.2026 – heute"     (pipe separator)
      "Title\t04.2026 – heute"      (tab separator)
      "Title (2023–heute)"          (parenthesised year range)
      "Title | 03.2026 – 04.2026"   (month-dot-year range)

    Returns None if the line does not end with a German date period.
    """
    m = _GERMAN_DATE_PERIOD_RE.search(line)
    if not m:
        return None
    date_str = m.group(0).strip()
    after = line[m.end():].strip()
    before = line[:m.start()]

    if after == ')':
        # Parenthesised date: strip the opening '(' from before
        before = re.sub(r'\s*\(\s*$', '', before)
    elif after:
        return None  # unexpected trailing content

    title = re.sub(r'[\s|·\t]+$', '', before).strip()
    if not title:
        return None
    return (title, date_str)


# ── Profile photo helper ──────────────────────────────────────────────────────

def _decode_photo(data_url: str) -> bytes | None:
    """
    Decode a data:image/*;base64,... URL to raw image bytes.
    Returns None if the URL is not a valid data URL or decoding fails.
    """
    m = re.match(r'^data:image/[^;]+;base64,(.+)$', data_url, re.DOTALL)
    if not m:
        return None
    try:
        return base64.b64decode(m.group(1))
    except Exception:
        return None


# ── CV DOCX ───────────────────────────────────────────────────────────────────

def generate_cv_docx(
    cv_markdown: str,
    candidate_name: str = "",
    candidate_city: str = "",
    candidate_email: str = "",
    candidate_phone: str = "",
    candidate_linkedin: str = "",
    candidate_github: str = "",
    candidate_photo_url: str = "",
) -> bytes:
    """
    Convert the tailored CV markdown to a DOCX file.
    Returns bytes suitable for a file download response.
    """
    doc = Document()
    _setup_a4(doc, top_cm=2.0, bottom_cm=2.0, left_cm=2.0, right_cm=2.0)
    _set_doc_default_font(doc, "Calibri", 10.5)

    # ── Deterministic header table ─────────────────────────────────────────────
    # Usable width = 21 - 2 - 2 = 17 cm.
    # Right column widens to 4 cm when a profile photo is available.
    usable_width_cm = 17.0

    photo_bytes = _decode_photo(candidate_photo_url) if candidate_photo_url else None
    right_col_cm = 4.0 if photo_bytes else 2.5
    left_col_cm = usable_width_cm - right_col_cm

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    _remove_table_borders(table)
    table.columns[0].width = Cm(left_col_cm)
    table.columns[1].width = Cm(right_col_cm)

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)

    for cell in (left_cell, right_cell):
        for para in cell.paragraphs:
            _set_para_spacing(para, before_pt=0, after_pt=0)

    # Name (20 pt, blue, bold)
    name_para = left_cell.paragraphs[0]
    _set_para_spacing(name_para, before_pt=0, after_pt=4)
    if candidate_name:
        run = name_para.add_run(candidate_name)
        _set_run_font(run, size_pt=20, color=_HEADING_COLOR, bold=True)

    # Contact line: city | phone | email
    # No application job-title line — the CV header shows only personal contact info.
    city_short = _short_city(candidate_city)
    contact_parts = [x for x in [city_short, candidate_phone, candidate_email] if x]
    if contact_parts:
        cp = left_cell.add_paragraph()
        _set_para_spacing(cp, before_pt=0, after_pt=2)
        _add_run(cp, "  |  ".join(contact_parts), size_pt=9.5, color=_DATE_COLOR)

    # Links line: LinkedIn | GitHub
    link_parts = []
    if candidate_linkedin:
        link_parts.append(_clean_url(candidate_linkedin))
    if candidate_github:
        link_parts.append(_clean_url(candidate_github))
    if link_parts:
        lp = left_cell.add_paragraph()
        _set_para_spacing(lp, before_pt=0, after_pt=2)
        _add_run(lp, "  |  ".join(link_parts), size_pt=9.5, color=_DATE_COLOR)

    # Right cell: portrait-cropped profile photo when available; empty otherwise.
    photo_para = right_cell.paragraphs[0]
    photo_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_para_spacing(photo_para, before_pt=0, after_pt=0)
    if photo_bytes:
        portrait_bytes = _crop_portrait(photo_bytes)
        if portrait_bytes:
            try:
                run = photo_para.add_run()
                # 225×300 px source (3:4) → width=3.0 cm, height auto-scales to 4.0 cm
                run.add_picture(io.BytesIO(portrait_bytes), width=Cm(3.0))
            except Exception:
                pass  # invalid image — right cell stays empty

    # ── Normalise and render markdown body ─────────────────────────────────────
    # TODO: exact 2-page enforcement requires a semantic/layout step — not in this patch.
    normalised = _normalize_cv_markdown(cv_markdown, candidate_name)
    seen_sections: set[str] = set()
    in_skip_section = False

    for line in normalised.splitlines():
        stripped = line.strip()

        # Section heading detection
        hm = _HEADING_RE.match(stripped)
        if hm:
            section_name = hm.group(2).strip().upper()
            if section_name in seen_sections:
                # Skip duplicate section and its content
                in_skip_section = True
            else:
                seen_sections.add(section_name)
                in_skip_section = False
                _add_cv_section_heading(doc, section_name)
            continue

        if in_skip_section or not stripped:
            continue

        if stripped.startswith(("- ", "• ", "* ")):
            _add_cv_bullet(doc, stripped[2:].strip())
        elif stripped.startswith("– ") or stripped.startswith("— "):
            _add_cv_bullet(doc, stripped[2:].strip())
        elif _TECH_STACK_LINE_RE.match(stripped):
            pass  # omit Tech-Stack lines; technologies are in KENNTNISSE
        else:
            entry = _split_title_date(stripped)
            if entry:
                title, date = entry
                _add_cv_entry_title(doc, title, period=date,
                                    usable_width_cm=usable_width_cm)
            else:
                para = doc.add_paragraph()
                _set_para_spacing(para, before_pt=1, after_pt=1)
                _render_inline(para, stripped, size_pt=10.5)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Anschreiben DOCX ──────────────────────────────────────────────────────────

_DE_MONTHS = ["Januar", "Februar", "März", "April", "Mai", "Juni",
               "Juli", "August", "September", "Oktober", "November", "Dezember"]


def generate_anschreiben_docx(
    anschreiben_text: str,
    candidate_name: str = "",
    candidate_street: str = "",
    candidate_postal_code: str = "",
    candidate_city: str = "",
    candidate_phone: str = "",
    candidate_email: str = "",
    candidate_linkedin: str = "",
    candidate_github: str = "",  # kept in signature but not shown in letter header
    company_name: str = "",
    company_address: str = "",
    candidate_address: str = "",  # legacy alias for candidate_city
) -> bytes:
    """
    Convert the Anschreiben to a DOCX following the German reference letter style.
    Returns bytes suitable for a file download response.
    """
    if candidate_address and not candidate_city:
        candidate_city = candidate_address

    doc = Document()
    _setup_a4(doc, top_cm=2.0, bottom_cm=2.0, left_cm=2.5, right_cm=2.0)
    _set_doc_default_font(doc, "Calibri", 11.0)

    # ── 1. Sender block (one field per line) ───────────────────────────────────
    if candidate_name:
        p = doc.add_paragraph()
        _set_para_spacing(p, before_pt=0, after_pt=4)
        _add_run(p, candidate_name, size_pt=11, bold=True)

    # Address line: "Burgstraße 20, 57072 Siegen" — or just city if no street/postal
    city_short = _short_city(candidate_city)
    postal_city = " ".join(x for x in [candidate_postal_code, city_short] if x)
    address_parts = [x for x in [candidate_street, postal_city] if x]
    if address_parts:
        p = doc.add_paragraph()
        _set_para_spacing(p, before_pt=0, after_pt=4)
        _add_run(p, ", ".join(address_parts), size_pt=11)
    elif city_short:
        p = doc.add_paragraph()
        _set_para_spacing(p, before_pt=0, after_pt=4)
        _add_run(p, city_short, size_pt=11)

    if candidate_phone:
        p = doc.add_paragraph()
        _set_para_spacing(p, before_pt=0, after_pt=4)
        _add_run(p, candidate_phone, size_pt=11)

    if candidate_email:
        p = doc.add_paragraph()
        _set_para_spacing(p, before_pt=0, after_pt=4)
        _add_run(p, candidate_email, size_pt=11)

    if candidate_linkedin:
        p = doc.add_paragraph()
        _set_para_spacing(p, before_pt=0, after_pt=4)
        _add_run(p, _clean_url(candidate_linkedin), size_pt=11)
    # Note: GitHub omitted from letter header — relevant for CV, not cover letter

    # ── 2. Blank line ──────────────────────────────────────────────────────────
    _add_blank_line(doc, size_pt=6)

    # ── 3. Recipient block ─────────────────────────────────────────────────────
    if company_name:
        p = doc.add_paragraph()
        _set_para_spacing(p, before_pt=0, after_pt=4)
        _add_run(p, company_name, size_pt=11, bold=True)
    # Company address: render only if explicitly provided — never invent it
    if company_address:
        for addr_line in company_address.splitlines():
            addr_line = addr_line.strip()
            if addr_line:
                p = doc.add_paragraph()
                _set_para_spacing(p, before_pt=0, after_pt=4)
                _add_run(p, addr_line, size_pt=11)

    # ── 4. Blank line ──────────────────────────────────────────────────────────
    _add_blank_line(doc, size_pt=6)

    # ── 5. Date, right-aligned ─────────────────────────────────────────────────
    _today = date.today()
    today_str = f"{_today.day}. {_DE_MONTHS[_today.month - 1]} {_today.year}"
    p = doc.add_paragraph()
    _set_para_spacing(p, before_pt=0, after_pt=8)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_run(p, today_str, size_pt=11)

    # ── 6–11. Body ─────────────────────────────────────────────────────────────
    _render_anschreiben_body(doc, anschreiben_text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _render_anschreiben_body(doc: Document, text: str) -> None:
    """
    Render Anschreiben body text into the document.

    Handles:
    - Subject line (bold): line starting with "Betreff:", "Bewerbung", "Ihre Ausschreibung"
    - Greeting: "Sehr geehrte..."
    - Body paragraphs: plain/inline-bold text
    - Closing: "Mit freundlichen Grüßen" — always gets a signature gap below it
    - Candidate name: lines immediately after closing (same or next blank-separated block)

    The closing + name are split even if no blank line separates them in the LLM output.
    """
    # Build blocks: lists of non-blank lines, separated by blank lines
    blocks: list[list[str]] = []
    current: list[str] = []
    for line in text.splitlines():
        s = line.strip()
        if s:
            current.append(s)
        else:
            if current:
                blocks.append(current)
                current = []
    if current:
        blocks.append(current)

    post_closing = False

    for block_lines in blocks:
        first = block_lines[0].lower()

        is_subject = (
            first.startswith("betreff:")
            or first.startswith("bewerbung")
            or first.startswith("ihre ausschreibung")
        )
        is_greeting = first.startswith("sehr geehrte")
        is_closing = first.startswith("mit freundlichen")

        if post_closing:
            # Lines after the closing block are the candidate name / signature
            for ln in block_lines:
                p = doc.add_paragraph()
                _set_para_spacing(p, before_pt=0, after_pt=4)
                _add_run(p, ln, size_pt=11)
            continue

        if is_closing:
            post_closing = True
            # Render the closing formula on its own paragraph
            p = doc.add_paragraph()
            _set_para_spacing(p, before_pt=12, after_pt=4)
            _add_run(p, block_lines[0], size_pt=11)
            # Signature gap (space for handwritten signature)
            for _ in range(2):
                _add_blank_line(doc, size_pt=8)
            # Any additional lines in the same block (name merged without blank line)
            for name_line in block_lines[1:]:
                if name_line.strip():
                    p = doc.add_paragraph()
                    _set_para_spacing(p, before_pt=0, after_pt=4)
                    _add_run(p, name_line, size_pt=11)
            continue

        # Join lines within the block for subject / greeting / body
        block_text = " ".join(block_lines)

        if is_subject:
            p = doc.add_paragraph()
            _set_para_spacing(p, before_pt=0, after_pt=8)
            subject_text = re.sub(r"^[Bb]etreff:\s*", "", block_text)
            _add_run(p, subject_text, size_pt=11, bold=True)
        elif is_greeting:
            p = doc.add_paragraph()
            _set_para_spacing(p, before_pt=0, after_pt=6)
            _add_run(p, block_text, size_pt=11)
        else:
            p = doc.add_paragraph()
            _set_para_spacing(p, before_pt=0, after_pt=6)
            _render_inline(p, block_text, size_pt=11)
